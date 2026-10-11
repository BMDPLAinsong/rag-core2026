"""Word 文档加载器（基于 python-docx）"""
from pathlib import Path
from typing import List

from docx import Document as DocxDocument

from core.models import Document
from loaders.base import BaseLoader


class DocxLoader(BaseLoader):
    def load(self, file_path: str) -> List[Document]:
        path = Path(file_path)
        doc = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            return []
        return [
            Document(
                source=path.name,
                text="\n".join(paragraphs),
                metadata={"file_type": "docx"},
            )
        ]
